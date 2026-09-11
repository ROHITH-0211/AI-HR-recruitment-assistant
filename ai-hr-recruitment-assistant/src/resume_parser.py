"""
Resume Parser
-------------
Extracts raw text from an uploaded PDF or DOCX resume, then asks the LLM
(via OpenRouter) to convert that text into clean, structured JSON data.
"""

import io

import fitz  # PyMuPDF
import docx

from src.llm import chat_json


RESUME_SCHEMA_INSTRUCTIONS = """
You are an expert resume parser for an HR recruitment system.

Read the resume text and extract the candidate's information into
STRICT JSON only (no markdown fences, no commentary, no extra text).

Use exactly this schema:
{
  "name": string,
  "email": string,
  "phone": string,
  "location": string,
  "summary": string,
  "education": [
    {"degree": string, "institution": string, "year": string, "score": string}
  ],
  "experience": [
    {"title": string, "company": string, "duration": string, "description": string}
  ],
  "job_history": [
    {"title": string, "company": string, "start_date": string, "end_date": string}
  ],
  "skills": [string],
  "certifications": [string],
  "projects": [
    {"name": string, "description": string, "technologies": [string]}
  ],
  "total_experience_years": number
}

Rules:
- If a field is missing from the resume, use an empty string, empty list, or 0.
- "skills" should be a flat, deduplicated list of technical and soft skills.
- "job_history" is a short chronological list distinct from "experience" (which holds fuller descriptions).
- Keep descriptions concise (1-2 sentences).
- Return ONLY the JSON object, nothing else.
"""


class ResumeParsingError(Exception):
    """Raised when a resume file can't be read or has no usable content."""


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract plain text from a PDF file's bytes using PyMuPDF."""
    try:
        text_parts = []
        with fitz.open(stream=file_bytes, filetype="pdf") as pdf:
            for page in pdf:
                text_parts.append(page.get_text())
        return "\n".join(text_parts).strip()
    except Exception as exc:
        raise ResumeParsingError(f"This PDF could not be read (it may be corrupted or scanned/image-only): {exc}")


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract plain text from a DOCX file's bytes using python-docx."""
    try:
        document = docx.Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text)
        return "\n".join(paragraphs).strip()
    except Exception as exc:
        raise ResumeParsingError(f"This DOCX file could not be read (it may be corrupted): {exc}")


def extract_resume_text(file_bytes: bytes, filename: str) -> str:
    """Route to the correct extractor based on file extension."""
    if not file_bytes:
        raise ResumeParsingError("The uploaded resume file is empty.")

    lower_name = filename.lower()
    if lower_name.endswith(".pdf"):
        text = extract_text_from_pdf(file_bytes)
    elif lower_name.endswith(".docx"):
        text = extract_text_from_docx(file_bytes)
    else:
        raise ResumeParsingError("Unsupported file type. Please upload a PDF or DOCX resume.")

    if not text.strip():
        raise ResumeParsingError(
            "No readable text was found in this resume. It may be a scanned image without a text layer."
        )
    return text


def parse_resume_with_llm(resume_text: str) -> dict:
    """Send extracted resume text to the LLM and get back structured JSON."""
    if not resume_text.strip():
        raise ResumeParsingError("No readable text was found in this resume file.")
    return chat_json(RESUME_SCHEMA_INSTRUCTIONS, f"Resume text:\n\n{resume_text[:12000]}")


def parse_resume(file_bytes: bytes, filename: str) -> dict:
    """Full pipeline: file bytes -> extracted text -> structured candidate data."""
    text = extract_resume_text(file_bytes, filename)
    structured = parse_resume_with_llm(text)
    structured["_raw_text"] = text
    return structured
